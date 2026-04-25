from pathlib import Path

from docx import Document

from app.ingestion.docx_ingestion_artifact import ingest_docx_file
from app.ingestion.normalized_artifact import build_normalized_artifact
from app.ingestion.table_chunker import chunk_tables_from_artifact


def test_chunk_tables_from_artifact_basic_case(tmp_path: Path) -> None:
    file_path = tmp_path / "FS_FCIS_14.7.0.0.0$ASNB_R24_Teller_Branch_Reports_Realignment_v1.0.docx"

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Column"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "Branch Code"
    table.cell(1, 1).text = "Unique branch identifier"
    document.save(file_path)

    raw_artifact = ingest_docx_file(file_path)
    normalized_artifact = build_normalized_artifact(raw_artifact)
    chunked_tables = chunk_tables_from_artifact(normalized_artifact)

    assert chunked_tables.document_family == "FS_FCIS_14.7.0.0.0$ASNB"
    assert chunked_tables.release_label == "R24"
    assert chunked_tables.total_table_chunks == 1
    assert chunked_tables.table_chunks[0].row_count == 2
    assert chunked_tables.table_chunks[0].text == (
        "Column | Description\n"
        "Branch Code | Unique branch identifier"
    )
